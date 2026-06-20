"""
Seed-document text extraction.

Converts an uploaded seed / knowledge document into plain text for storage in
`projects.document_text` and downstream LLM profile extraction. Supports PDF (pypdf),
DOCX (python-docx), and plain-text formats (TXT / MD / JSON / CSV / …).

The browser can only read plain text via FileReader, so binary formats (PDF / DOCX)
are sent to the backend and extracted here.
"""

from __future__ import annotations
import io


def extract_file_text(filename: str, data: bytes) -> str:
    """Return plain text extracted from an uploaded file's bytes.

    Detection is by extension first, with a PDF magic-byte fallback. Unknown / text types
    are decoded as UTF-8 (errors ignored). Returns "" when nothing could be extracted.
    """
    if not data:
        return ""
    name = (filename or "").lower()

    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        return _extract_pdf(data)
    if name.endswith(".docx"):
        return _extract_docx(data)
    # Plain-text formats (txt / md / json / csv / …) — decode tolerantly.
    return data.decode("utf-8", errors="ignore")


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
        return "\n\n".join(p for p in parts if p).strip()
    except Exception as e:
        print(f"[FileText] PDF extraction failed ({e})")
        return ""


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
        # Include table cell text — specs often put fields in tables.
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()
    except Exception as e:
        print(f"[FileText] DOCX extraction failed ({e})")
        return ""
